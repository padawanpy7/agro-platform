"""Escribe el Excel de estimacion a partir del plan (JSON) que arma scripts/estimacion.js.

No inventa el formato: PARTE DE LA PLANTILLA (scripts/lib/estimacion-plantilla.xlsx, que es el
"Estimacion - ICC-91.xlsx" real que se importa) y solo reemplaza contenido. Asi las columnas,
los anchos, los estilos y los formatos de fecha quedan identicos al archivo que espera el
importador: si se armara la planilla de cero, cualquier diferencia de estilo o de cabecera
rompe la importacion y no se nota hasta que falla.

Uso:  python estimacion-xlsx.py <plan.json> <plantilla.xlsx> <salida.xlsx> [--valores]
      --valores escribe numeros ya calculados en vez de formulas (para importadores que leen
      el valor cacheado: un .xlsx recien generado tiene la formula pero no su resultado).
"""

import json
import sys
from copy import copy
from datetime import date

FILA_DATOS = 6  # primera fila de tareas en la plantilla
FILAS_PLANTILLA = 5  # cuantas filas de tareas trae la plantilla (6..10)
COLS = "ABCDEFGH"


def _fecha(iso):
    a, m, d = (int(x) for x in iso.split("-"))
    return date(a, m, d)


def _clonar_estilo(ws, fila_origen, fila_destino):
    for col in range(1, len(COLS) + 2):
        origen = ws.cell(row=fila_origen, column=col)
        destino = ws.cell(row=fila_destino, column=col)
        destino._style = copy(origen._style)
    ws.row_dimensions[fila_destino].height = ws.row_dimensions[fila_origen].height


def escribir(plan, plantilla, salida, valores=False):
    import openpyxl

    wb = openpyxl.load_workbook(plantilla)
    ws = wb.active
    ws.title = plan["hoja"][:31]

    ws["B1"] = plan["jira"]
    ws["B2"] = plan["proyecto"]
    ws["E3"] = plan["minutosPorPeso"]
    ws["B4"] = plan["minutosJornada"]

    filas = plan["filas"]
    ultima = FILA_DATOS + len(filas) - 1

    if len(filas) > FILAS_PLANTILLA:
        for f in range(FILA_DATOS + FILAS_PLANTILLA, ultima + 1):
            _clonar_estilo(ws, FILA_DATOS, f)
    elif len(filas) < FILAS_PLANTILLA:
        ws.delete_rows(ultima + 1, FILAS_PLANTILLA - len(filas))
        for f in range(ultima + 1, FILA_DATOS + FILAS_PLANTILLA):
            ws.row_dimensions.pop(f, None)

    for i, f in enumerate(filas):
        r = FILA_DATOS + i
        ws.cell(row=r, column=1).value = f["tarea"]
        ws.cell(row=r, column=2).value = f["descripcion"]
        ws.cell(row=r, column=3).value = f["peso"]
        ws.cell(row=r, column=4).value = f["minutos"] if valores else f"=C{r}*$E$3"
        ws.cell(row=r, column=5).value = f["horas"] if valores else f"=D{r}/60"
        for col, clave in ((6, "inicio"), (7, "fin"), (8, "entrega")):
            c = ws.cell(row=r, column=col)
            c.value = _fecha(f[clave])
            c.number_format = "d-mmm"

    t = plan["totales"]
    ws["F3"] = t["minutos"] if valores else f"=SUM(D{FILA_DATOS}:D{ultima})"
    ws["G3"] = t["horas"] if valores else "=F3/60"
    ws["H3"] = t["horasPorDia"] if valores else f"=G3/{plan['divisorDias']}"
    import math

    ws["I3"] = math.ceil(t["horasPorDia"]) if valores else "=ROUNDUP(H3,0)"

    wb.save(salida)


def main(argv):
    args = [a for a in argv[1:] if not a.startswith("--")]
    if len(args) != 3:
        print(__doc__, file=sys.stderr)
        return 2
    plan_json, plantilla, salida = args
    with open(plan_json, encoding="utf-8") as fh:
        plan = json.load(fh)
    escribir(plan, plantilla, salida, valores="--valores" in argv[1:])
    print(salida)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
